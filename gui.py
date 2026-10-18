import requests
from urllib.parse import urlencode, urlparse
import iv8
import os
import json
import datetime
import tkinter as tk
from tkinter import ttk, scrolledtext, messagebox, filedialog
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import threading

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BDMS_JS_PATH = os.path.join(SCRIPT_DIR, "bdms_1.0.1.19_fix.js")
COOKIES_PATH = os.path.join(SCRIPT_DIR, "cookies.json")

with open(BDMS_JS_PATH, 'r', encoding='utf-8') as f:
    BDMS_JS = f.read()

DEFAULT_ENVIRONMENT = {
    "location": {
        "href": "https://www.douyin.com/video/7682671884655596901",
        "origin": "https://www.douyin.com",
        "protocol": "https:", "host": "www.douyin.com",
        "hostname": "www.douyin.com", "port": "",
        "pathname": "/video/7682671884655596901",
        "search": "", "hash": ""
    },
    "navigator": {
        "userAgent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/145.0.0.0 Safari/537.36",
        "language": "zh-CN", "languages": ["zh-CN", "zh", "en"],
    },
}

cookies_data = {}
if os.path.exists(COOKIES_PATH):
    with open(COOKIES_PATH, 'r', encoding='utf-8') as f:
        cookies_data = json.load(f)

headers = {
    "authority": "www.douyin.com",
    "accept": "application/json, text/plain, */*",
    "referer": "https://www.douyin.com/",
    "user-agent": DEFAULT_ENVIRONMENT["navigator"]["userAgent"],
}

_last_comments_data = []
_last_raw_data = None
_current_aweme_id = ""

def generate_abogus(base_url, params):
    with iv8.JSContext(environment=DEFAULT_ENVIRONMENT) as ctx:
        ctx.eval("""window.MessageChannel = __iv8__.wrapNative(function() {
            var a={onmessage:null}, b={onmessage:null};
            a.postMessage = function(d) { if(b.onmessage) setTimeout(function(){b.onmessage({data:d})}, 0); };
            b.postMessage = function(d) { if(a.onmessage) setTimeout(function(){a.onmessage({data:d})}, 0); };
            return {port1:a, port2:b};
        }, "MessageChannel");""")
        ctx.eval(BDMS_JS)
        init_params = '{"aid":6383,"pageId":6241,"paths":["^/aweme/v1/"],"boe":false,"ddrt":8.5,"ic":8.5}'
        ctx.eval('window.bdms.init(' + init_params + ');')
        query_string = urlencode(params, safe='*')
        request_url = base_url + "?" + query_string
        result = ctx.eval('var xhr=new XMLHttpRequest();xhr.open("GET","' + request_url + '",true);xhr.send(null);window.__iv8__.netLog.entries;', to_py=True)
        if result and len(result) > 0:
            return result[0]["url"]
        raise RuntimeError("iv8 未能生成签名")

def fetch_comments(aweme_id, cursor, count):
    base_url = "https://www.douyin.com/aweme/v1/web/comment/list/"
    params = {
        "device_platform": "webapp", "aid": "6383",
        "channel": "channel_pc_web", "aweme_id": aweme_id,
        "cursor": str(cursor), "count": str(count),
        "item_type": "0", "insert_ids": "",
        "whale_cut_token": "", "cut_version": "1",
        "rcFT": "", "update_version_code": "170400",
        "pc_client_type": "1", "pc_libra_divert": "Windows",
        "support_h265": "1", "support_dash": "1",
        "cpu_core_num": "16", "version_code": "170400",
        "version_name": "17.4.0", "cookie_enabled": "true",
        "screen_width": "1707", "screen_height": "960",
        "browser_language": "zh-CN", "browser_platform": "Win32",
        "browser_name": "Chrome", "browser_version": "145.0.0.0",
        "browser_online": "true", "engine_name": "Blink",
        "engine_version": "145.0.0.0", "os_name": "Windows",
        "os_version": "10", "device_memory": "8", "platform": "PC",
        "downlink": "2.7", "effective_type": "4g",
        "round_trip_time": "100",
    }
    full_url = generate_abogus(base_url, params)
    resp = requests.get(full_url, headers=headers, cookies=cookies_data, timeout=15)
    if resp.status_code != 200:
        raise RuntimeError(f"HTTP {resp.status_code}: {resp.text[:200]}")
    raw = resp.text.strip()
    if not raw:
        raise RuntimeError("API 返回空响应，请检查 cookies 是否过期")
    data = resp.json()
    if not isinstance(data, dict):
        raise RuntimeError(f"API 返回非 JSON 数据: {raw[:100]}")
    return data

def _build_verify_params():
    """reply 接口额外的会话校验参数：verifyFp/fp(=s_v_web_id) + webid + uifid"""
    sv = cookies_data.get("s_v_web_id", "")
    webid = cookies_data.get("webid", "")
    uifid = cookies_data.get("UIFID", "")
    p = {}
    if sv:
        p["verifyFp"] = sv
        p["fp"] = sv
    if webid:
        p["webid"] = webid
    if uifid:
        p["uifid"] = uifid
    return p

def fetch_replies(aweme_id, reply_id, cursor, count):
    base_url = "https://www.douyin.com/aweme/v1/web/comment/list/reply/"
    params = {
        "device_platform": "webapp", "aid": "6383",
        "channel": "channel_pc_web",
        "item_id": aweme_id,
        "comment_id": str(reply_id),
        "cut_version": "1",
        "cursor": str(cursor), "count": str(count),
        "item_type": "0",
        "update_version_code": "170400",
        "pc_client_type": "1", "pc_libra_divert": "Windows",
        "support_h265": "1", "support_dash": "1",
        "cpu_core_num": "16", "version_code": "170400",
        "version_name": "17.4.0", "cookie_enabled": "true",
        "screen_width": "1707", "screen_height": "960",
        "browser_language": "zh-CN", "browser_platform": "Win32",
        "browser_name": "Chrome", "browser_version": "145.0.0.0",
        "browser_online": "true", "engine_name": "Blink",
        "engine_version": "145.0.0.0", "os_name": "Windows",
        "os_version": "10", "device_memory": "8", "platform": "PC",
        "downlink": "10", "effective_type": "4g",
        "round_trip_time": "150",
    }
    params.update(_build_verify_params())
    full_url = generate_abogus(base_url, params)
    resp = requests.get(full_url, headers=headers, cookies=cookies_data, timeout=15)
    if resp.status_code != 200:
        return []
    raw = resp.text.strip()
    if not raw:
        return []
    data = resp.json()
    return data.get("comments") or []

def load_reply_comments(aweme_id, comment_widgets):
    """后台线程加载所有有子评论的评论的子评论"""
    global _last_comments_data
    threads = []
    results = {}

    def fetch_one(cid, total):
        limit = min(int(total), 50)
        all_replies = []
        cursor = 0
        while cursor < limit:
            batch = min(20, limit - cursor)
            reps = fetch_replies(aweme_id, cid, cursor, batch)
            if not reps:
                break
            all_replies.extend(reps)
            if len(reps) < batch:
                break
            cursor += batch
        results[cid] = all_replies

    for cid, total in comment_widgets.items():
        t = threading.Thread(target=fetch_one, args=(cid, total), daemon=True)
        t.start()
        threads.append(t)

    for t in threads:
        t.join()

    return results

def do_fetch():
    global _last_comments_data, _last_raw_data, _current_aweme_id
    aweme_id = entry_aweme.get().strip()
    page_str = entry_page.get().strip()
    count_str = entry_count.get().strip()

    if not aweme_id:
        messagebox.showwarning("提示", "请输入视频 ID")
        return
    try:
        page = int(page_str)
    except ValueError:
        messagebox.showwarning("提示", "页码必须是数字")
        return
    try:
        count = int(count_str)
    except ValueError:
        messagebox.showwarning("提示", "每页条数必须是数字")
        return

    cursor = page * count
    btn.config(state="disabled")

    now = datetime.datetime.now().strftime("%H:%M:%S")
    txt_log.config(state="normal")
    txt_log.delete("1.0", tk.END)
    txt_log.insert(tk.END, f"[{now}] 请求中... aweme_id={aweme_id} cursor={cursor}\n")
    txt_log.see(tk.END)
    txt_log.config(state="disabled")
    root.update()

    comments_data = None
    try:
        _last_raw_data = data = fetch_comments(aweme_id, cursor, count)
        _current_aweme_id = aweme_id

        status_code = data.get("status_code")
        all_comments = data.get("comments") or []
        has_more = data.get("has_more", 0)
        next_cursor = data.get("cursor", 0)
        total = len(all_comments)

        txt_log.config(state="normal")
        txt_log.insert(tk.END, f"[API status_code={status_code}] 本帧 {total} 条评论  has_more={has_more} next_cursor={next_cursor}\n")
        txt_log.see(tk.END)

        _last_comments_data.clear()
        tree.delete(*tree.get_children())

        comment_widgets = {}

        for c in all_comments:
            user = c.get("user") or {}
            nickname = str(user.get("nickname", "?"))
            uid = str(user.get("uid", ""))
            text = str(c.get("text") or "").replace("\n", " ")
            digg = int(c.get("digg_count") or 0)
            replies_count = int(c.get("reply_comment_total") or 0)
            create_time = int(c.get("create_time") or 0)
            ip_label = str(c.get("ip_label") or "")
            ct = datetime.datetime.fromtimestamp(create_time).strftime("%Y-%m-%d %H:%M") if create_time else ""
            cid = c.get("cid", "")

            tree.insert("", "end", values=(nickname, uid, text[:80], digg, replies_count, ip_label, ct, ""), tags=("parent",))
            _last_comments_data.append({"level": 0, "昵称": nickname, "UID": uid, "评论内容": text, "点赞数": digg, "回复数": replies_count, "IP属地": ip_label, "发布时间": ct, "回复对象": ""})

            if replies_count > 0 and cid:
                comment_widgets[cid] = replies_count

        # txt_log.insert(tk.END, f"\n主评论已显示 {_last_comments_data.__len__()} 条  有子评论数: {len(comment_widgets)}\n")
        txt_log.see(tk.END)

        if comment_widgets:
            # txt_log.insert(tk.END, f"[{datetime.datetime.now().strftime('%H:%M:%S')}] 正在后台加载子评论...\n")
            txt_log.see(tk.END)
            txt_log.config(state="disabled")
            root.update()

            def load_and_refresh():
                results = load_reply_comments(_current_aweme_id, comment_widgets)
                tree.delete(*tree.get_children())
                _last_comments_data.clear()
                loaded_count = 0
                for c in all_comments:
                    user = c.get("user") or {}
                    nickname = str(user.get("nickname", "?"))
                    uid = str(user.get("uid", ""))
                    text = str(c.get("text") or "").replace("\n", " ")
                    digg = int(c.get("digg_count") or 0)
                    replies_count = int(c.get("reply_comment_total") or 0)
                    create_time = int(c.get("create_time") or 0)
                    ip_label = str(c.get("ip_label") or "")
                    ct = datetime.datetime.fromtimestamp(create_time).strftime("%Y-%m-%d %H:%M") if create_time else ""
                    cid = c.get("cid", "")

                    parent_item = tree.insert("", "end", values=(nickname, uid, text[:80], digg, replies_count, ip_label, ct, ""), tags=("parent",))
                    _last_comments_data.append({"level": 0, "昵称": nickname, "UID": uid, "评论内容": text, "点赞数": digg, "回复数": replies_count, "IP属地": ip_label, "发布时间": ct, "回复对象": ""})

                    replies = results.get(cid, [])
                    for rc in replies:
                        ru = rc.get("user") or {}
                        rnickname = str(ru.get("nickname", "?"))
                        ruid = str(ru.get("uid", ""))
                        rtext = str(rc.get("text") or "").replace("\n", " ")
                        rdigg = int(rc.get("digg_count") or 0)
                        rreplies = int(rc.get("reply_comment_total") or 0)
                        rtime = int(rc.get("create_time") or 0)
                        rip = str(rc.get("ip_label") or "")
                        rct = datetime.datetime.fromtimestamp(rtime).strftime("%Y-%m-%d %H:%M") if rtime else ""
                        tree.insert(parent_item, "end", values=(rnickname, ruid, rtext[:80], rdigg, rreplies, rip, rct, "@" + nickname), tags=("child",))
                        _last_comments_data.append({"level": 1, "昵称": rnickname, "UID": ruid, "评论内容": rtext, "点赞数": rdigg, "回复数": rreplies, "IP属地": rip, "发布时间": rct, "回复对象": nickname})
                        loaded_count += 1

                    tree.item(parent_item, open=True)

                txt_log.config(state="normal")
                # txt_log.insert(tk.END, f"[{datetime.datetime.now().strftime('%H:%M:%S')}] 子评论加载完成，共加载 {loaded_count} 条\n")
                txt_log.see(tk.END)
                txt_log.config(state="disabled")
                btn.config(state="normal")

            t = threading.Thread(target=load_and_refresh, daemon=True)
            t.start()
        else:
            txt_log.insert(tk.END, f"[{datetime.datetime.now().strftime('%H:%M:%S')}] 无子评论，下一页 cursor={next_cursor} (页码={page+1})\n")
            txt_log.see(tk.END)
            txt_log.config(state="disabled")

    except Exception as e:
        import traceback
        err_msg = f"[ERROR] {e}\n{traceback.format_exc()}"
        txt_log.config(state="normal")
        txt_log.insert(tk.END, err_msg + "\n")
        txt_log.see(tk.END)
        txt_log.config(state="disabled")
        messagebox.showerror("请求失败", f"{e}")
    finally:
        if not comment_widgets:
            btn.config(state="normal")

def do_save_excel():
    if not _last_comments_data:
        messagebox.showwarning("提示", "还没有抓取数据，请先点击「抓取」")
        return
    now_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    aweme_id = entry_aweme.get().strip() or "video"
    default_name = f"抖音评论_{aweme_id}_{now_str}.xlsx"
    path = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")],
        initialfile=default_name,
        initialdir=SCRIPT_DIR,
    )
    if not path:
        return
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "评论数据"
        col_headers = ["层级", "序号", "昵称", "UID", "评论内容", "点赞数", "回复数", "IP属地", "发布时间", "回复对象"]
        header_fill = PatternFill("solid", fgColor="0078D4")
        header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
        thin = Side(style="thin", color="BBBBBB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for col, h in enumerate(col_headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        row_idx = 2
        for idx, item in enumerate(_last_comments_data):
            level_text = "主评论" if item.get("level") == 0 else "子评论"
            ws.cell(row=row_idx, column=1, value=level_text).border = border
            ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=idx+1).border = border
            ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=3, value=item["昵称"]).border = border
            ws.cell(row=row_idx, column=4, value=str(item["UID"])).border = border
            ws.cell(row=row_idx, column=5, value=item["评论内容"]).border = border
            ws.cell(row=row_idx, column=6, value=item["点赞数"]).border = border
            ws.cell(row=row_idx, column=6).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=7, value=item["回复数"]).border = border
            ws.cell(row=row_idx, column=7).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=8, value=item["IP属地"]).border = border
            ws.cell(row=row_idx, column=8).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=9, value=item["发布时间"]).border = border
            ws.cell(row=row_idx, column=9).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=10, value=item.get("回复对象", "")).border = border
            ws.cell(row=row_idx, column=10).alignment = Alignment(horizontal="center")
            if row_idx % 2 == 0:
                for col in range(1, 11):
                    ws.cell(row=row_idx, column=col).fill = PatternFill("solid", fgColor="F2F2F2")
            row_idx += 1
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 8
        ws.column_dimensions["C"].width = 14
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 60
        ws.column_dimensions["F"].width = 10
        ws.column_dimensions["G"].width = 10
        ws.column_dimensions["H"].width = 10
        ws.column_dimensions["I"].width = 18
        ws.column_dimensions["J"].width = 14
        for row in range(2, len(_last_comments_data) + 2):
            ws.cell(row=row, column=5).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[1].height = 24
        for row in range(2, len(_last_comments_data) + 2):
            ws.row_dimensions[row].height = 36
        wb.save(path)
        messagebox.showinfo("成功", f"已保存 {len(_last_comments_data)} 条评论\n{path}")
        txt_log.config(state="normal")
        txt_log.insert(tk.END, f"[{datetime.datetime.now().strftime('%H:%M:%S')}] 已保存 Excel: {os.path.basename(path)}\n")
        txt_log.see(tk.END)
        txt_log.config(state="disabled")
    except Exception as e:
        messagebox.showerror("保存失败", str(e))

root = tk.Tk()
root.title("抖音评论抓取器 (a_bogus)")
root.geometry("1100x700")
root.configure(bg="white")

frame_top = tk.Frame(root, bg="white")
frame_top.pack(fill="x", padx=16, pady=12)

tk.Label(frame_top, text="视频 ID:", bg="white", fg="#333", font=("Microsoft YaHei", 10)).grid(row=0, column=0, sticky="w", padx=(0, 6))
entry_aweme = tk.Entry(frame_top, width=36, font=("Consolas", 10), bg="#f5f5f5", fg="#222", insertbackground="#000", relief="solid", bd=1)
entry_aweme.grid(row=0, column=1, padx=4)
entry_aweme.insert(0, "7682671884655596901")
tk.Label(frame_top, text="页码(0=第一页):", bg="white", fg="#333", font=("Microsoft YaHei", 10)).grid(row=0, column=2, sticky="w", padx=(24, 6))
entry_page = tk.Entry(frame_top, width=8, font=("Consolas", 10), bg="#f5f5f5", fg="#222", insertbackground="#000", relief="solid", bd=1)
entry_page.grid(row=0, column=3, padx=4)
entry_page.insert(0, "0")
tk.Label(frame_top, text="每页条数:", bg="white", fg="#333", font=("Microsoft YaHei", 10)).grid(row=0, column=4, sticky="w", padx=(12, 6))
entry_count = tk.Entry(frame_top, width=6, font=("Consolas", 10), bg="#f5f5f5", fg="#222", insertbackground="#000", relief="solid", bd=1)
entry_count.grid(row=0, column=5, padx=4)
entry_count.insert(0, "10")
btn = tk.Button(frame_top, text="  抓 取  ", font=("Microsoft YaHei", 10, "bold"), bg="#0078d4", fg="white",
                activebackground="#1084d8", activeforeground="white", relief="flat",
                padx=24, pady=6, cursor="hand2", command=do_fetch)
btn.grid(row=0, column=6, padx=(20, 0))
btn_save = tk.Button(frame_top, text=" 保存Excel ", font=("Microsoft YaHei", 10), bg="#388e3c", fg="white",
                      activebackground="#4caf50", activeforeground="white", relief="flat",
                      padx=18, pady=6, cursor="hand2", command=do_save_excel)
btn_save.grid(row=0, column=7, padx=(10, 0))

frame_table = tk.Frame(root, bg="white")
frame_table.pack(fill="both", expand=True, padx=16, pady=(0, 8))

columns = ("nickname", "uid", "content", "likes", "replies", "ip", "time", "reply_to")
tree = ttk.Treeview(frame_table, columns=columns, show="headings", selectmode="extended")
tree.tag_configure("parent")
tree.tag_configure("child")
tree.heading("nickname", text="用户名")
tree.heading("uid", text="UID")
tree.heading("content", text="评论内容")
tree.heading("likes", text="点赞数")
tree.heading("replies", text="回复数")
tree.heading("ip", text="IP属地")
tree.heading("time", text="发布时间")
tree.heading("reply_to", text="回复")
tree.column("nickname", width=120, minwidth=80, anchor="w")
tree.column("uid", width=130, minwidth=100, anchor="w")
tree.column("content", width=320, minwidth=160, anchor="w")
tree.column("likes", width=65, minwidth=50, anchor="center")
tree.column("replies", width=65, minwidth=50, anchor="center")
tree.column("ip", width=75, minwidth=60, anchor="center")
tree.column("time", width=135, minwidth=100, anchor="center")
tree.column("reply_to", width=80, minwidth=60, anchor="w")
tree.pack(side="left", fill="both", expand=True)

sb = tk.Scrollbar(frame_table, orient="vertical", command=tree.yview)
sb.pack(side="right", fill="y")
tree.config(yscrollcommand=sb.set)

def on_double_click(event):
    item = tree.selection()
    if not item:
        return
    vals = tree.item(item, "values")
    # 兼容旧版本（含类型列，9 列）与新版本（8 列）
    if len(vals) == 9:
        flag, nickname, uid, content, likes, replies, ip, time, reply_to = vals
    else:
        nickname, uid, content, likes, replies, ip, time, reply_to = vals
        flag = ""
    info = f"UID: {uid}   IP: {ip}   {time}   赞:{likes}  回复:{replies}"
    top = tk.Toplevel(root)
    top.title(f"@{nickname} 的评论")
    top.configure(bg="white")
    top.geometry("560x380")
    top.transient(root)
    top.grab_set()
    tk.Label(top, text=f"@{nickname}", font=("Microsoft YaHei", 12, "bold"), bg="white", fg="#0078d4").pack(pady=(10, 4))
    info = f"UID: {uid}   IP: {ip}   {time}   赞:{likes}  回复:{replies}"
    tk.Label(top, text=info, font=("Consolas", 9), bg="white", fg="#666").pack(pady=(0, 8))
    txt = scrolledtext.ScrolledText(top, font=("Microsoft YaHei", 10), bg="white", fg="#222", wrap="word", relief="flat")
    txt.pack(fill="both", expand=True, padx=12, pady=(0, 8))
    txt.insert("1.0", content)
    if reply_to:
        txt.insert("1.0", f"回复 @{reply_to}: ")
    txt.config(state="disabled")
    btn_close = tk.Button(top, text="关闭", font=("Microsoft YaHei", 9), bg="#eee", fg="#333",
                          relief="flat", padx=16, pady=3, cursor="hand2", command=top.destroy)
    btn_close.pack(pady=(0, 10))

if __name__ == '__main__':
    tree.bind("<Double-1>", on_double_click)

    frame_log = tk.Frame(root, bg="white")
    frame_log.pack(fill="x", padx=16, pady=(0, 10))

    txt_log = scrolledtext.ScrolledText(frame_log, font=("Consolas", 9), bg="#fafafa", fg="#444",
                                         insertbackground="#000", wrap="word", relief="solid", bd=1, height=5)
    txt_log.pack(fill="x")
    txt_log.config(state="disabled")
    txt_log.insert(tk.END, "日志区 — 请求完成后显示状态信息\n")

    status_bar = tk.Label(root, text=f"cookies: {len(cookies_data)} 个 | bdms: bdms_1.0.1.19_fix.js",
                           bg="white", fg="#888", font=("Microsoft YaHei", 9), anchor="w")
    status_bar.pack(fill="x")

    root.mainloop()
